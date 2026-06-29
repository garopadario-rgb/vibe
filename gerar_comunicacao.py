from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Comunicação & Apresentação"

# ── Paleta ─────────────────────────────────────────────────────────────────
DARK_BG      = "1B1F3B"
HEADER_COL   = "2D3561"
WHITE        = "FFFFFF"
GRAY_TEXT    = "555577"

# Categoria — Comunicação/Apresentação Geral
CAT1_DARK    = "0A3D62"
CAT1_LIGHT   = "D6EAF8"
CAT1_ALT     = "EBF5FB"
CAT1_ACCENT  = "1A5276"

# Categoria — YouTube / Câmera
CAT2_DARK    = "4A235A"
CAT2_LIGHT   = "F5EEF8"
CAT2_ALT     = "FAF5FF"
CAT2_ACCENT  = "7D3C98"

thin   = Side(style="thin",   color="CCCCCC")
medium = Side(style="medium", color="888888")
b_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
b_med  = Border(left=medium, right=medium, top=medium, bottom=medium)

# ── Título principal ───────────────────────────────────────────────────────
ws.merge_cells("A1:J1")
c = ws["A1"]
c.value = "🎤 20 MELHORES CURSOS: COMUNICAÇÃO & TÉCNICAS DE APRESENTAÇÃO"
c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 38

ws.merge_cells("A2:J2")
s = ws["A2"]
s.value = "15 cursos de comunicação & apresentação  •  5 cursos focados em YouTube / câmera  |  Todos os formatos: PT & EN"
s.font = Font(name="Calibri", size=10, italic=True, color="AAAACC")
s.fill = PatternFill("solid", fgColor=DARK_BG)
s.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

ws.row_dimensions[3].height = 6   # espaço

# ── Cabeçalhos ─────────────────────────────────────────────────────────────
headers = ["#", "Título do Vídeo/Curso", "Link YouTube", "Categoria",
           "Subtema", "Canal / Criador", "Língua", "Data", "Visualizações", "Notas"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col, value=h)
    cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=HEADER_COL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = b_med
ws.row_dimensions[4].height = 24

# ── Dados ──────────────────────────────────────────────────────────────────
data = [
    # ── BLOCO 1: Separador ──
    # ── COMUNICAÇÃO & APRESENTAÇÃO GERAL (1–15) ───────────────────────────
    (1,
     "Complete Presentation Skills Course: Master Storytelling, Slide Design & Public Speaking",
     "https://www.youtube.com/watch?v=NJlBdiTUL-8",
     "Comunicação Geral", "Storytelling + Slides + Oratória",
     "Verificar no YT", "Inglês", "Mai 2025", "Verificar no YT",
     "Curso completo e gratuito — storytelling, design de slides e falar em público"),

    (2,
     "Public Speaking for Beginners – Free Full Course | Speak with Confidence",
     "https://www.youtube.com/watch?v=MXKkXXYUxBc",
     "Comunicação Geral", "Oratória para Iniciantes",
     "Verificar no YT", "Inglês", "Set 2025", "Verificar no YT",
     "Curso completo de oratória do zero — confiança e presença em público"),

    (3,
     "Mastering Effective Presentations | Full Course | Presentation Skills",
     "https://www.youtube.com/watch?v=aRIQgQxGpss",
     "Comunicação Geral", "Apresentações de Impacto",
     "Verificar no YT", "Inglês", "Dez 2023", "Verificar no YT",
     "Masterclass completa — criar e entregar apresentações memoráveis"),

    (4,
     "Mastering Public Speaking & Presentation Skills: Full Course | Free Certificate 🎤",
     "https://www.youtube.com/watch?v=vgUUGM65xFs",
     "Comunicação Geral", "Falar em Público + Certificado",
     "Verificar no YT", "Inglês", "Dez 2025", "Verificar no YT",
     "Curso com certificado gratuito — oratória e apresentação profissional"),

    (5,
     "Free Presentation Course – Develop the Skills to Create & Deliver Amazing Presentations",
     "https://www.youtube.com/watch?v=gHXSwOTuseA",
     "Comunicação Geral", "Criar & Entregar Apresentações",
     "Verificar no YT", "Inglês", "2023", "Verificar no YT",
     "Toolkit completo para criar e apresentar com impacto — tudo revelado"),

    (6,
     "How to Present with Impact and Confidence | Master Public Speaking & Communication Skills",
     "https://www.youtube.com/watch?v=u1xRlsrfPyU",
     "Comunicação Geral", "Impacto & Confiança",
     "Verificar no YT", "Inglês", "Jan 2025", "Verificar no YT",
     "Técnicas de impacto e confiança para qualquer tipo de apresentação"),

    (7,
     "Communication Skills Course: How to Read Body Language and More — FREE COURSE",
     "https://www.youtube.com/watch?v=vfXQ18m_OsE",
     "Comunicação Geral", "Linguagem Corporal",
     "Verificar no YT", "Inglês", "2022", "Verificar no YT",
     "Curso gratuito — leitura de linguagem corporal e comunicação não-verbal"),

    (8,
     "Mastering the Art of Body Language: The Ultimate Non-Verbal Communication Course",
     "https://www.youtube.com/watch?v=fPb1Ersu0pk",
     "Comunicação Geral", "Comunicação Não-Verbal",
     "Verificar no YT", "Inglês", "2023", "Verificar no YT",
     "Dominar a linguagem corporal — curso completo e prático"),

    (9,
     "Free Masterclass: Confident Communication for a Better 2025",
     "https://www.youtube.com/watch?v=fntKY1Y0xGI",
     "Comunicação Geral", "Comunicação Confiante",
     "Verificar no YT", "Inglês", "Jan 2025", "Verificar no YT",
     "3 passos para comunicar com confiança em qualquer contexto"),

    (10,
     "Curso COMPLETO de Oratória para INICIANTES",
     "https://www.youtube.com/watch?v=JkhA2cxeOaA",
     "Comunicação Geral", "Oratória — Iniciantes",
     "Verificar no YT", "Português (BR)", "Mar 2025", "Verificar no YT",
     "Curso completo em PT — tudo o que um iniciante precisa saber"),

    (11,
     "Curso de Oratória Completo — Comunicação de Alta Performance #01",
     "https://www.youtube.com/watch?v=cTQHrNOlAUo",
     "Comunicação Geral", "Comunicação de Alta Performance",
     "Verificar no YT", "Português (BR)", "Jul 2021", "Verificar no YT",
     "Superar o medo de falar em público — estratégias de alta performance"),

    (12,
     "CURSO DE ORATÓRIA: Aprenda a Falar em Público com Confiança",
     "https://www.youtube.com/watch?v=y8N8W0Q3RIk",
     "Comunicação Geral", "Falar em Público",
     "Verificar no YT", "Português (BR)", "Abr 2024", "Verificar no YT",
     "Método completo para falar em público sem medo — passo a passo em PT"),

    (13,
     "10 Técnicas de Comunicação Assertiva que Todo Profissional Precisa Conhecer",
     "https://www.youtube.com/watch?v=2LABVbAgStA",
     "Comunicação Geral", "Comunicação Assertiva",
     "Verificar no YT", "Português (BR)", "Nov 2024", "Verificar no YT",
     "10 técnicas práticas para clareza, impacto e assertividade"),

    (14,
     "Aprenda a Falar em Público com Boa Didática, de Forma Clara e Objectiva!",
     "https://www.youtube.com/watch?v=WwaNX9HhJ_8",
     "Comunicação Geral", "Didáctica & Clareza",
     "Verificar no YT", "Português (BR)", "Mar 2025", "Verificar no YT",
     "Foco em didáctica — transmitir ideias com clareza e objectividade"),

    (15,
     "Oratória Persuasiva — Os Segredos da Comunicação para Convencer e Influenciar",
     "https://www.youtube.com/watch?v=N05mgOlRfFg",
     "Comunicação Geral", "Persuasão & Influência",
     "Verificar no YT", "Português (BR)", "2023", "Verificar no YT",
     "Técnicas de persuasão e influência para comunicadores profissionais"),

    # ── YOUTUBE & CÂMERA (16–20) ──────────────────────────────────────────
    (16,
     "Talk to the Camera Like a Pro (Even If You're Shy) — YouTube Creator Tips",
     "https://www.youtube.com/watch?v=tW3WDNjT6gk",
     "YouTube & Câmera", "Presença em Câmera",
     "Verificar no YT", "Inglês", "Jul 2025", "Verificar no YT",
     "Técnicas para falar para a câmera com confiança — ideal para youtubers"),

    (17,
     "How to TALK to CAMERA! (Complete Course) — Be Natural & Engaging on Video",
     "https://www.youtube.com/watch?v=Z8wI6QvJEdk",
     "YouTube & Câmera", "Curso Câmera Completo",
     "Verificar no YT", "Inglês", "2023", "Verificar no YT",
     "Curso completo de presença em câmera — naturalidade e envolvimento"),

    (18,
     "Tell Stories So Good You Finally Fix Your Sh*tty Retention (Full Course for YouTubers)",
     "https://www.youtube.com/watch?v=epKEXCHjp4M",
     "YouTube & Câmera", "Storytelling para YouTube",
     "Verificar no YT", "Inglês", "Set 2025", "Verificar no YT",
     "Storytelling para reter audiência no YouTube — workbook gratuito incluído"),

    (19,
     "Como Falar com Confiança na Frente das Câmeras | Dicas de Postura e Presença",
     "https://www.youtube.com/watch?v=JW8uwCbFCS0",
     "YouTube & Câmera", "Postura & Presença em Câmera",
     "Verificar no YT", "Português (BR)", "Out 2025", "Verificar no YT",
     "Postura, presença e confiança diante das câmeras — dicas em Português"),

    (20,
     "5 SEGREDOS PARA FALAR DIANTE DAS CÂMERAS — Falar para Câmera #1",
     "https://www.youtube.com/watch?v=vhGsIAG1Boc",
     "YouTube & Câmera", "Segredos da Câmera",
     "Verificar no YT", "Português (BR)", "2022", "Verificar no YT",
     "5 segredos práticos para criar vídeos melhores e monetizar o canal"),
]

for i, row in enumerate(data):
    er = i + 5
    is_yt = row[3] == "YouTube & Câmera"
    is_alt = i % 2 == 1

    if is_yt:
        bg = CAT2_ALT if is_alt else CAT2_LIGHT
        cat_color = CAT2_ACCENT
    else:
        bg = CAT1_ALT if is_alt else CAT1_LIGHT
        cat_color = CAT1_ACCENT

    fill = PatternFill("solid", fgColor=bg)

    for col, val in enumerate(row, 1):
        cell = ws.cell(row=er, column=col, value=val)
        cell.fill = fill
        cell.border = b_thin
        cell.alignment = Alignment(
            vertical="center", wrap_text=True,
            horizontal="center" if col in (1, 4, 5, 7, 8, 9) else "left"
        )

        if col == 1:
            cell.font = Font(name="Calibri", size=11, bold=True, color=DARK_BG)
        elif col == 2:
            cell.font = Font(name="Calibri", size=9, bold=True, color="1A1A2E")
        elif col == 3:
            cell.font = Font(name="Calibri", size=8, color="1155CC", underline="single")
            cell.hyperlink = val
            cell.value = val
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        elif col == 4:
            cell.font = Font(name="Calibri", size=9, bold=True, color=cat_color)
        else:
            cell.font = Font(name="Calibri", size=9, color="2C2C2C")

    ws.row_dimensions[er].height = 38

# ── Larguras ───────────────────────────────────────────────────────────────
widths = {1:4, 2:46, 3:16, 4:18, 5:22, 6:22, 7:12, 8:10, 9:16, 10:44}
for col, w in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# ── Rodapé ─────────────────────────────────────────────────────────────────
fr = len(data) + 6
ws.merge_cells(f"A{fr}:J{fr}")
leg = ws[f"A{fr}"]
leg.value = (
    "🔵 Azul = Comunicação & Apresentação Geral (15 cursos)  |  "
    "🟣 Roxo = YouTube & Câmera (5 cursos especializados)  |  "
    "⚠️ 'Verificar no YT' = abrir o link e anotar visualizações actuais  |  Dados: Junho 2026"
)
leg.font = Font(name="Calibri", size=8, italic=True, color=GRAY_TEXT)
leg.fill = PatternFill("solid", fgColor="EEEEFF")
leg.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[fr].height = 28

ws.freeze_panes = "A5"

path = "/home/user/vibe/Comunicacao_Apresentacao_YouTube.xlsx"
wb.save(path)
print(f"Guardado: {path}")
