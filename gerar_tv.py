from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Comunicação Televisiva"

# ── Paleta ─────────────────────────────────────────────────────────────────
DARK_BG    = "1A0A00"   # marrom escuro TV
HEADER_COL = "8B0000"   # vermelho profundo (estudio)
WHITE      = "FFFFFF"
GRAY_TEXT  = "555555"

# Categoria EN
EN_LIGHT   = "FDECEA"
EN_ALT     = "FFF5F4"
EN_ACCENT  = "C0392B"

# Categoria PT
PT_LIGHT   = "FEF9E7"
PT_ALT     = "FFFDF0"
PT_ACCENT  = "B7770D"

thin   = Side(style="thin",   color="CCCCCC")
medium = Side(style="medium", color="999999")
b_thin = Border(left=thin,   right=thin,   top=thin,   bottom=thin)
b_med  = Border(left=medium, right=medium, top=medium, bottom=medium)

# ── Título ─────────────────────────────────────────────────────────────────
ws.merge_cells("A1:J1")
c = ws["A1"]
c.value = "📺 20 MELHORES CURSOS: COMUNICAÇÃO & APRESENTAÇÃO TELEVISIVA"
c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 38

ws.merge_cells("A2:J2")
s = ws["A2"]
s.value = (
    "Técnicas de apresentação TV  •  Âncoras & Jornalismo Broadcast  •  "
    "Media Training  •  Teleprompter  •  Locução  |  EN & PT"
)
s.font = Font(name="Calibri", size=10, italic=True, color="DDAA88")
s.fill = PatternFill("solid", fgColor=DARK_BG)
s.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 6

# ── Cabeçalhos ─────────────────────────────────────────────────────────────
headers = ["#", "Título do Vídeo/Curso", "Link YouTube", "Subtema",
           "Nível", "Canal / Criador", "Língua", "Data", "Visualizações", "Notas"]
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col, value=h)
    cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=HEADER_COL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = b_med
ws.row_dimensions[4].height = 24

# ── Dados ──────────────────────────────────────────────────────────────────
data = [
    # ── INGLÊS (1–15) ─────────────────────────────────────────────────────
    (1,
     "TV Presenting Masterclass — How to Become a TV Presenter: Full Step-by-Step",
     "https://www.youtube.com/watch?v=FnOiVmAdbNY",
     "Masterclass Apresentação TV", "Iniciante → Avançado",
     "Presenter Academy", "Inglês", "2022", "Verificar no YT",
     "Curso completo: direção passo a passo, dicas de apresentação e como conseguir trabalho na TV"),

    (2,
     "TV Presenter Training Pt1 — On-Camera Technique & Confidence",
     "https://www.youtube.com/watch?v=FqLF5YBsk8E",
     "Treino Apresentador TV", "Iniciante",
     "Presenter Academy / Brian Naylor", "Inglês", "2021", "Verificar no YT",
     "Parte 1: técnica em câmera, postura, contacto visual e confiança em directo"),

    (3,
     "TV Presenter Training Pt2 — Delivery, Tone & Presenting to Camera",
     "https://www.youtube.com/watch?v=AVCk_wgK8nk",
     "Treino Apresentador TV", "Intermédio",
     "Presenter Academy / Brian Naylor", "Inglês", "2021", "Verificar no YT",
     "Parte 2: entrega, tom de voz, ritmo e apresentação natural para câmera"),

    (4,
     "How to Become a TV Presenter — 4 Essential Steps (Webinar Highlights)",
     "https://www.youtube.com/watch?v=nlnxGdsNo5I",
     "Carreira em TV", "Iniciante → Intermédio",
     "Brian Naylor", "Inglês", "2022", "Verificar no YT",
     "4 passos essenciais para iniciar carreira como apresentador televisivo profissional"),

    (5,
     "Presenting Skills to Become a Great On-Screen Communicator",
     "https://www.youtube.com/watch?v=-_VzPT8ZSYw",
     "Comunicação em Ecrã", "Intermédio",
     "Verificar no YT", "Inglês", "2023", "Verificar no YT",
     "Habilidades de apresentação para comunicar com autoridade e naturalidade no ecrã"),

    (6,
     "Broadcast Journalism 101 — Fundamentals for On-Camera Journalists",
     "https://www.youtube.com/watch?v=zLBDaojhK78",
     "Jornalismo Broadcast", "Iniciante",
     "NBCU Academy", "Inglês", "Jun 2025", "Verificar no YT",
     "Fundamentos do jornalismo televisivo para quem quer aparecer em câmera"),

    (7,
     "Fundamentals of Journalism — Free Course Online (NBC News Group)",
     "https://www.youtube.com/watch?v=OzpU9QgmD8o",
     "Jornalismo Televisivo", "Iniciante",
     "NBCUniversal / NBCU Academy", "Inglês", "2023", "Verificar no YT",
     "Curso gratuito da NBC: fundamentos do jornalismo, reportagem e narrativa televisiva"),

    (8,
     "Anchoring and Teleprompters — NBCU Academy 101 (with MSNBC Anchor Yasmin Vossoughian)",
     "https://www.youtube.com/watch?v=O3BqluVM_7c",
     "Ancoragem & Teleprompter", "Intermédio",
     "NBCU Academy / MSNBC", "Inglês", "2022", "Verificar no YT",
     "Técnicas de ancoragem e uso do teleprompter com âncora profissional da MSNBC"),

    (9,
     "8 Simple Tips to Use a Teleprompter Like a Pro | Media Training",
     "https://www.youtube.com/watch?v=ZKfLjHw0Lq4",
     "Teleprompter & Leitura", "Intermédio",
     "Verificar no YT", "Inglês", "2021", "Verificar no YT",
     "Dominar o teleprompter — ler com naturalidade como os grandes âncoras televisivos"),

    (10,
     "How to Sound Natural On Air — How to Read a Teleprompter | Tips for Anchors & Reporters",
     "https://www.youtube.com/watch?v=82ToIEdbf8Y",
     "Naturalidade em Directo", "Intermédio → Avançado",
     "Verificar no YT", "Inglês", "Out 2025", "Verificar no YT",
     "Soar natural em directo — técnicas para âncoras e repórteres TV"),

    (11,
     "Improve Your Life by Speaking Like a News Anchor — Master Class",
     "https://www.youtube.com/watch?v=bSLAfuXz5pc",
     "Dicção & Voz Profissional", "Iniciante → Avançado",
     "Verificar no YT", "Inglês", "2023", "Verificar no YT",
     "Masterclass: adoptar a dicção, ritmo e autoridade vocal dos grandes âncoras de TV"),

    (12,
     "SPEAK Like a PRO News Anchor with These 5 Easy Tips!",
     "https://www.youtube.com/watch?v=is3l-ntx_Eo",
     "Voz & Dicção TV", "Iniciante",
     "Verificar no YT", "Inglês", "Ago 2025", "Verificar no YT",
     "5 dicas práticas para falar com autoridade e clareza como âncora profissional"),

    (13,
     "Media Training Tips: How to Present Well On Camera (for TV & Entrevistas)",
     "https://www.youtube.com/watch?v=2zQ0TXT7mfo",
     "Media Training", "Intermédio",
     "Verificar no YT", "Inglês", "Fev 2024", "Verificar no YT",
     "Treino de media para entrevistas televisivas — comunicação clara e de impacto"),

    (14,
     "Media Training Tips 101: What You Need to Know to Be Like a Pro",
     "https://www.youtube.com/watch?v=NWVNduCBCJI",
     "Media Training", "Iniciante",
     "Verificar no YT", "Inglês", "2023", "Verificar no YT",
     "Media training completo: como dominar entrevistas e aparecer na televisão como profissional"),

    (15,
     "TV Presenter Tips: 5 Starter Tips to Be a Good TV Host (Beginners)",
     "https://www.youtube.com/watch?v=NF9TRNR47DI",
     "Anfitrião TV", "Iniciante",
     "Presenter Academy", "Inglês", "2022", "Verificar no YT",
     "5 dicas essenciais para iniciantes: postura, ritmo, contacto visual e autoridade televisiva"),

    # ── PORTUGUÊS (16–20) ─────────────────────────────────────────────────
    (16,
     "Curso Gratuito de Apresentador de Televisão — Técnicas Profissionais",
     "https://www.youtube.com/watch?v=4YECNw345Co",
     "Apresentador TV", "Iniciante",
     "Verificar no YT", "Português (BR)", "2022", "Verificar no YT",
     "Curso gratuito completo: técnicas profissionais de apresentação televisiva em PT"),

    (17,
     "Curso para Apresentador de TV, Locutor de Rádio e Narrador — Desinibição",
     "https://www.youtube.com/watch?v=zo2qtSuN38g",
     "Locução & Narração TV", "Iniciante → Intermédio",
     "Delphis Fonseca", "Português (BR)", "2021", "Verificar no YT",
     "Desenvolver voz, dicção e desinibição para televisão, rádio e narração"),

    (18,
     "Como Ser Um Apresentador de TV — Técnicas e Carreira",
     "https://www.youtube.com/watch?v=nleoPRDDl40",
     "Carreira & Técnicas TV", "Iniciante",
     "Watson Weber", "Português (BR)", "2022", "Verificar no YT",
     "Dicas práticas de carreira e técnicas de comunicação televisiva em Português"),

    (19,
     "AULA DE APRESENTAÇÃO DE PROGRAMA DE TV | Comradio — Técnicas ao Vivo",
     "https://www.youtube.com/watch?v=5_Ymv85zd64",
     "Apresentação de Programa TV", "Intermédio",
     "Comradio", "Português (BR)", "2021", "Verificar no YT",
     "Aula prática de apresentação de programa televisivo — técnicas ao vivo em estúdio"),

    (20,
     "ESTUDANTES DO CURSO DE APRESENTAÇÃO DE TV E REPÓRTAGEM — Prática em Estúdio",
     "https://www.youtube.com/watch?v=qaXKDTG5-FM",
     "Repórtagem & Apresentação", "Iniciante → Intermédio",
     "TV Porcelana / Curso Prático", "Português (BR)", "Abr 2024", "Verificar no YT",
     "Aulas práticas de repórtagem e apresentação televisiva com certificado — filmagem real em estúdio"),
]

for i, row in enumerate(data):
    er = i + 5
    is_pt = "Português" in row[6]
    is_alt = i % 2 == 1

    if is_pt:
        bg = PT_ALT if is_alt else PT_LIGHT
        acc = PT_ACCENT
    else:
        bg = EN_ALT if is_alt else EN_LIGHT
        acc = EN_ACCENT

    fill = PatternFill("solid", fgColor=bg)

    for col, val in enumerate(row, 1):
        cell = ws.cell(row=er, column=col, value=val)
        cell.fill = fill
        cell.border = b_thin
        cell.alignment = Alignment(
            vertical="center", wrap_text=True,
            horizontal="center" if col in (1, 4, 5, 7, 8, 9) else "left"
        )
        if col == 1:
            cell.font = Font(name="Calibri", size=11, bold=True, color=DARK_BG)
        elif col == 2:
            cell.font = Font(name="Calibri", size=9, bold=True, color="1A1A2E")
        elif col == 3:
            cell.font = Font(name="Calibri", size=8, color="1155CC", underline="single")
            cell.hyperlink = val
            cell.value = val
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        elif col == 4:
            cell.font = Font(name="Calibri", size=9, bold=True, color=acc)
        else:
            cell.font = Font(name="Calibri", size=9, color="2C2C2C")

    ws.row_dimensions[er].height = 40

# ── Larguras ───────────────────────────────────────────────────────────────
widths = {1:4, 2:46, 3:38, 4:22, 5:20, 6:22, 7:12, 8:10, 9:16, 10:46}
for col, w in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# ── Rodapé ─────────────────────────────────────────────────────────────────
fr = len(data) + 6
ws.merge_cells(f"A{fr}:J{fr}")
leg = ws[f"A{fr}"]
leg.value = (
    "🔴 Vermelho = Cursos em Inglês (15 vídeos)  |  🟡 Amarelo = Cursos em Português (5 vídeos)  |  "
    "Links clicáveis directamente na coluna 'Link YouTube'  |  Dados: Junho 2026"
)
leg.font = Font(name="Calibri", size=8, italic=True, color=GRAY_TEXT)
leg.fill = PatternFill("solid", fgColor="FFF8F0")
leg.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[fr].height = 24

ws.freeze_panes = "A5"

path = "/home/user/vibe/Comunicacao_Televisiva.xlsx"
wb.save(path)
print(f"Guardado: {path}")
